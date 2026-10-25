"""Build native editable Word content from MinerU 4 middle_json and reviewed config."""
from pathlib import Path
import argparse,io,json,re,hashlib
from PIL import Image
from lxml import html as LH
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from common import textof,plain,children,location
from content_guard import table_hazards,validate_review

class Builder:
 def __init__(self,data,pages,cfg):
  self.data=data;self.pages=pages;self.cfg=cfg;self.doc=Document()
  self.report={'schema_version':2,'status':'draft_requires_final_acceptance','headings':[],'tables':0,'images':0,'joined_tables':[],'skipped':[],'review_items':[],'source_blocks':0,'table_decisions':[],'cell_images':[],'code_blocks':[]}
  self.source_reviews={}
  for item in cfg.get('_reviews',{}).get('table_reviews',[]):
   key=(int(item['page']),int(item['block_index']))
   if key in self.source_reviews:raise ValueError('Duplicate source review location')
   self.source_reviews[key]=item
  self.last_kind=None;self.last_table=None;self.last_cols=0;self.last_location=None
  self.page=0;self.block=0
  self.joins={(int(x['page']),int(x['block_index'])) for x in cfg.get('table_joins',[])}
  self.used_joins=set()
  self.overrides={(int(x['page']),int(x['block_index'])):int(x['level']) for x in cfg.get('heading_overrides',[])}
  self.ignore={(int(x['page']),int(x['block_index'])):x.get('reason','') for x in cfg.get('ignore_blocks',[])}
  self.excluded=set(cfg.get('exclude_pages',[]));self.width=float(cfg.get('page_width_cm',21))-2*float(cfg.get('margin_cm',2))
  if self.width<=0:raise ValueError('Invalid page width or margins')
  self.expected={}
  for page in data['pages']:
   for block in page.get('blocks',[]):
    if block['type']=='index':
     for entry in children(block):
      m=re.match(r'^(\d+(?:[.．]\d+)*)\s*(.+?)(?:\.{2,}|…|·{2,})',plain(entry))
      if m:self.expected[m[1].replace('．','.')]=re.sub(r'^[.．]\s*','',m[2].strip())
  self.setup()

 def review(self,reason,**details):
  self.report['review_items'].append({'page':self.page,'block_index':self.block,'reason':reason,**details})

 def clean(self,value):
  text=plain(value)
  for fix in self.cfg.get('text_replacements',[]):
   if int(fix['page'])==self.page and fix['old'] in text:text=text.replace(fix['old'],fix['new'])
  if re.search(r'\\[a-zA-Z]+|\$[^$]+\$',text):self.review('possible_unconverted_formula',text=text[:180])
  return text

 def setup(self):
  doc=self.doc;sec=doc.sections[0];sec.page_width=Cm(float(self.cfg.get('page_width_cm',21)));sec.page_height=Cm(float(self.cfg.get('page_height_cm',29.7)))
  sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(float(self.cfg.get('margin_cm',2)))
  sec.different_first_page_header_footer=bool(self.cfg.get('cover'))
  for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Heading 4','Heading 5','TOC Heading']:
   style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(float(self.cfg.get('body_font_pt',11)))
   style.font.color.rgb=RGBColor(0,0,0);style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
   style.paragraph_format.space_after=Pt(5);style.paragraph_format.line_spacing=1.25
   if style._element.pPr is not None:
    for x in list(style._element.pPr):
     if x.tag==qn('w:pBdr'):style._element.pPr.remove(x)
  doc.styles['Normal'].paragraph_format.first_line_indent=Pt(22)
  # Remove theme font fallbacks that can silently select Japanese fonts.
  for style in doc.styles:
   fonts=style._element.get_or_add_rPr().get_or_add_rFonts()
   for attr in list(fonts.attrib):
    if attr.lower().endswith('theme'):del fonts.attrib[attr]
   fonts.set(qn('w:eastAsia'),self.cfg.get('east_asia_font','宋体'))
  for name,size in [('Title',22),('Heading 1',16),('Heading 2',14),('Heading 3',12),('Heading 4',11),('Heading 5',11)]:
   style=doc.styles[name];style.font.size=Pt(size);style.font.bold=True
   style.paragraph_format.space_before=Pt(12);style.paragraph_format.first_line_indent=Pt(0);style.paragraph_format.keep_with_next=True
  for level in (1,2,3,4,5):
   name='TOC '+str(level)
   try:style=doc.styles[name]
   except KeyError:style=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
   style.font.name='Times New Roman';style.font.size=Pt(11);style.font.color.rgb=RGBColor(0,0,0)
   style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
   style.paragraph_format.first_line_indent=Pt(0);style.paragraph_format.left_indent=Cm((level-1)*.65)
   style.paragraph_format.space_after=Pt(4);style.paragraph_format.line_spacing=1.15
  p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;self.field(p,'PAGE')
  update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings._element.append(update)

 @staticmethod
 def field(p,instruction):
  run=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');run._r.append(begin)
  text=OxmlElement('w:instrText');text.set(qn('xml:space'),'preserve');text.text=' '+instruction+' ';run._r.append(text)
  end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');run._r.append(end)

 def para(self,value,heading=False,bold=False,title=False):
  text=self.clean(value)
  if not text:return
  level=self.overrides.get((self.page,self.block));m=re.match(r'^(\d+(?:[.．]\d+){0,5})[.．]?\s*([^\d].*)$',text)
  if m:
   number=m[1].replace('．','.');wanted=self.expected.get(number)
   if wanted and re.sub(r'\s','',wanted)==re.sub(r'\s','',m[2]):level=min(number.count('.')+1,5)
   elif heading and len(text)<100:level=level or min(number.count('.')+1,5)
   if level:text=number+' '+m[2].strip()
  if level and not 1<=level<=5:raise ValueError('Heading level must be 1..5')
  p=self.doc.add_paragraph(text,style='Title' if title else 'Heading '+str(level) if level else 'Normal')
  self.body_start(p)
  if level:self.report['headings'].append({'page':self.page,'block_index':self.block,'level':level,'text':text})
  elif heading or bold:
   p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
   for run in p.runs:run.bold=True
  elif re.match(r'^[（(]?\d+[）).、]|^[◆◇✧\-•]',text):p.paragraph_format.first_line_indent=Pt(0)
  self.last_kind='text';self.last_table=None;self.last_location=(self.page,self.block)
  return p

 def body_start(self,p):
  if getattr(self,'body_page_break_pending',False):
   p.paragraph_format.page_break_before=True
   self.body_page_break_pending=False

 def code(self,block):
  # Code is literal content: prose normalization would alter CJK strings,
  # escape sequences, entities and the first line's indentation.
  body=next((x for x in children(block) if x.get('type')=='code_body'),block)
  text=textof(body).replace('\r\n','\n').replace('\r','\n')
  for fix in self.cfg.get('text_replacements',[]):
   if int(fix['page'])==self.page and fix['old'] in text:text=text.replace(fix['old'],fix['new'])
  if not text:return
  p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
  self.body_start(p)
  p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.05
  run=p.add_run(text);run.font.name='Consolas';run.font.size=Pt(9)
  run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),self.cfg.get('east_asia_font','宋体'))
  self.report['code_blocks'].append({'page':self.page,'block_index':self.block,'lines':text.count('\n')+1})
  self.review('code_requires_source_check')
  self.last_kind='code';self.last_table=None;self.last_location=(self.page,self.block)
  return p

 def table(self,markup,join=False,cell_images=None):
  if getattr(self,'body_page_break_pending',False):
   p=self.doc.add_paragraph();self.body_start(p)
   p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run().font.size=Pt(1)
  tree=LH.fromstring(markup);rows=tree.xpath('.//tr') if tree.tag!='tr' else [tree]
  if not rows:raise ValueError(f'Empty table at {self.page}/{self.block}')
  cells=[];occupied=set();nc=0
  for ri,row in enumerate(rows):
   ci=0
   for td in row.xpath('./td|./th'):
    while (ri,ci) in occupied:ci+=1
    rs=max(1,int(td.get('rowspan','1')));cs=max(1,int(td.get('colspan','1')))
    if ri+rs>len(rows):raise ValueError(f'Invalid rowspan at {self.page}/{self.block}')
    for rr in range(ri,ri+rs):
     for cc in range(ci,ci+cs):
      if (rr,cc) in occupied:raise ValueError('Overlapping HTML cells')
      occupied.add((rr,cc))
    for br in td.xpath('.//br'):br.tail='\n'+(br.tail or '')
    cells.append((ri,ci,rs,cs,self.clean(td.text_content())));ci+=cs;nc=max(nc,ci)
  if not nc or nc>24:raise ValueError('Unusual table grid; review instead of flattening to text')
  if join and not(self.last_kind=='table' and self.last_table is not None and nc==self.last_cols):
   raise ValueError(f'Confirmed join has incompatible predecessor at {self.page}/{self.block}')
  if self.last_kind=='table' and not join:
   p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(0)
   p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=Pt(1);p.add_run().font.size=Pt(1)
  table=self.doc.add_table(rows=len(rows),cols=nc);table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
  for ri,ci,rs,cs,text in cells:
   cell=table.cell(ri,ci)
   if rs>1 or cs>1:cell=cell.merge(table.cell(ri+rs-1,ci+cs-1))
   cell.text=text;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in cell.paragraphs:
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.1
    if ri==0 and not join:p.paragraph_format.keep_with_next=True
    for run in p.runs:run.font.size=Pt(9 if nc<8 else 8)
   if ri==0 and not join:
    for p in cell.paragraphs:
     for run in p.runs:run.bold=True
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');cell._tc.get_or_add_tcPr().append(shade)
  if not join:table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for image in cell_images or []:
   ri=int(image['row']);ci=int(image['col'])
   if not(0<=ri<len(rows) and 0<=ci<nc):raise ValueError('Cell image row/column is outside table grid')
   if not any(r==ri and c==ci for r,c,rs,cs,t in cells):raise ValueError('Cell image must target the top-left anchor of a logical merged cell')
   source_page=int(image.get('source_page',self.page));box=image['bbox']
   if not(1<=source_page<=len(self.pages) and len(box)==4 and 0<=box[0]<box[2]<=1 and 0<=box[1]<box[3]<=1):raise ValueError('Invalid cell image source crop')
   if not image.get('evidence'):raise ValueError('Cell image requires source evidence')
   with Image.open(self.pages[source_page-1]['path']) as source:
    w,h=source.size;crop=source.crop(tuple(round(v*(w if k%2==0 else h)) for k,v in enumerate(box)))
   stream=io.BytesIO();crop.save(stream,format='PNG');stream.seek(0)
   target=table.cell(ri,ci);p=target.add_paragraph() if target.text.strip() or target._tc.xpath('.//w:drawing') else target.paragraphs[0]
   p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   width=min(self.width/nc*max(1,int(target._tc.tcPr.gridSpan.val) if target._tc.tcPr.gridSpan is not None else 1)-.15,5*crop.width/crop.height)
   p.add_run().add_picture(stream,width=Cm(max(.1,width)));self.report['images']+=1
   self.report['cell_images'].append({'page':self.page,'block_index':self.block,'row':ri,'col':ci,'source_page':source_page,'bbox':box,'evidence':image['evidence'],
    'output_table_index':self.report['tables'] if join else self.report['tables']+1,
    'output_row':len(self.last_table.rows)+ri if join else ri})
  if join:
   origin=self.last_location
   for row in list(table._tbl.findall(qn('w:tr'))):self.last_table._tbl.append(row)
   table._tbl.getparent().remove(table._tbl);table=self.last_table
   self.report['joined_tables'].append({'page':self.page,'block_index':self.block,'previous':origin})
   self.used_joins.add((self.page,self.block))
   self.review('joined_table_seam_requires_visual_check')
  else:self.report['tables']+=1
  self.last_kind='table';self.last_table=table;self.last_cols=nc;self.last_location=(self.page,self.block)

 def picture(self,block):
  body=next((x for x in children(block) if x.get('type') in ('image_body','chart_body','figure_body')),block)
  box=body.get('bbox') or block.get('bbox')
  if not box:self.review('image_missing_bbox');return
  source=self.pages[self.page-1]
  with Image.open(source['path']) as opened:im=opened.convert('RGB')
  w,h=im.size;x1,y1,x2,y2=box
  if max(box)>2:x1/=w;x2/=w;y1/=h;y2/=h
  if not (0<=x1<x2<=1 and 0<=y1<y2<=1):raise ValueError(f'Invalid bbox at {self.page}/{self.block}')
  if (x2-x1)*(y2-y1)<float(self.cfg.get('min_figure_area',0)):
   self.report['skipped'].append({'page':self.page,'block_index':self.block,'reason':'configured_small_figure_threshold','bbox':box});return
  crop=im.crop((max(0,int(x1*w)-3),max(0,int(y1*h)-3),min(w,int(x2*w)+3),min(h,int(y2*h)+3)))
  width=min(self.width,self.width*(x2-x1)/.8,float(self.cfg.get('figure_max_height_cm',20.5))*crop.width/crop.height)
  stream=io.BytesIO();crop.save(stream,format='PNG');stream.seek(0)
  p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1
  self.body_start(p)
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(stream,width=Cm(width))
  self.report['images']+=1;self.last_kind='image';self.last_table=None

 def cover(self):
  cfg=self.cfg.get('cover')
  if not cfg:return
  page=int(cfg['source_page'])
  if page not in self.excluded:raise ValueError('Rebuilt cover source must be excluded to avoid duplication')
  box=cfg.get('logo_bbox')
  if box:
   with Image.open(self.pages[page-1]['path']) as opened:im=opened.convert('RGB')
   w,h=im.size
   if not (len(box)==4 and 0<=box[0]<box[2]<=1 and 0<=box[1]<box[3]<=1):raise ValueError('Invalid cover logo bbox')
   stream=io.BytesIO();im.crop(tuple(round(v*(w if i%2==0 else h)) for i,v in enumerate(box))).save(stream,format='PNG');stream.seek(0)
   p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(stream,width=Cm(min(7,self.width)))
   self.report['images']+=1
  p=self.doc.add_paragraph(self.cfg['title'],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Cm(5)
  if self.cfg.get('company'):
   p=self.doc.add_paragraph(self.cfg['company']);p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Cm(6)
  self.doc.add_page_break()

 def normalize_tables(self):
  twips=round(self.width/2.54*1440)
  for table in self.doc.tables:
   grid=table._tbl.tblGrid;nc=len(grid.gridCol_lst);lengths=[0]*nc
   for tr in table._tbl.tr_lst:
    ci=0
    for tc in tr.tc_lst:
     span=int(tc.tcPr.gridSpan.val) if tc.tcPr is not None and tc.tcPr.gridSpan is not None else 1
     value=''.join(tc.xpath('.//w:t/text()'))
     if span==1 and ci<nc:lengths[ci]=max(lengths[ci],min(len(value),100))
     ci+=span
    if ci!=nc:raise ValueError('Inconsistent Word table grid')
   weights=[max(8,min(35,x**.6*2.5)) for x in lengths];total=sum(weights)
   widths=[round(twips*x/total) for x in weights];widths[-1]+=twips-sum(widths)
   for col,width in zip(grid.gridCol_lst,widths):col.set(qn('w:w'),str(width))
   tw=table._tbl.tblPr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(twips));table.autofit=False
   for tr in table._tbl.tr_lst:
    ci=0
    for tc in tr.tc_lst:
     props=tc.get_or_add_tcPr();span=int(props.gridSpan.val) if props.gridSpan is not None else 1
     props.get_or_add_tcW().set(qn('w:w'),str(sum(widths[ci:ci+span])));props.tcW.set(qn('w:type'),'dxa');ci+=span
     max_cx=max(1,round((int(props.tcW.w)-120)*635))
     for drawing in tc.xpath('.//w:drawing'):
      for extent in drawing.xpath('.//wp:extent'):
       cx=int(extent.get('cx'));cy=int(extent.get('cy'))
       if cx>max_cx:
        new_cy=round(cy*max_cx/cx);extent.set('cx',str(max_cx));extent.set('cy',str(new_cy))
        for size in drawing.xpath('.//a:xfrm/a:ext'):size.set('cx',str(max_cx));size.set('cy',str(new_cy))

 def run(self):
  self.cover();toc_page=self.cfg.get('toc_before_page');inserted=False
  known_text={'text','paragraph_title','doc_title','code','list','ordered_list','unordered_list','equation','equation_inline','equation_interline'}
  for page in self.data['pages']:
   self.page=int(page['page_idx'])+1
   if self.page in self.excluded:
    self.report['skipped'].append({'page':self.page,'reason':'configured_excluded_page'});continue
   if self.page==toc_page:
    if self.doc.paragraphs or self.doc.tables:self.doc.add_page_break()
    p=self.doc.add_paragraph('目录','TOC Heading');p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
    levels=int(self.cfg.get('toc_levels',3))
    if not 1<=levels<=5:raise ValueError('toc_levels must be 1..5')
    self.field(p,f'TOC \\o "1-{levels}" \\h \\z \\u');self.body_page_break_pending=True;inserted=True;self.last_kind=None;self.last_table=None
   for block in page.get('blocks',[]):
    self.block=int(block['index']);key=(self.page,self.block);typ=block['type'];self.report['source_blocks']+=1
    if key in self.ignore:
     if typ=='table':raise ValueError('Use a reviewed table decision to omit table artifacts, not ignore_blocks')
     self.report['skipped'].append({'page':self.page,'block_index':self.block,'reason':self.ignore[key]});continue
    if typ in ('header','footer','page_number','index'):
     self.report['skipped'].append({'page':self.page,'block_index':self.block,'reason':'structural_'+typ});continue
    if typ=='table':
     bodies=[x for x in children(block) if x.get('type')=='table_body']
     if len(bodies)!=1:raise ValueError('Table must have exactly one body; inspect structure')
     review=self.source_reviews.get(key);decision=validate_review(review,bodies[0]['content'])
     self.report['table_decisions'].append({**review,'hazards':table_hazards(bodies[0]['content'])})
     if decision=='omit_artifact':
      if key in self.joins:raise ValueError('Cannot join an omitted artifact')
      self.report['skipped'].append({'page':self.page,'block_index':self.block,'reason':'source_verified_artifact','evidence':review['evidence']});continue
     for sub in children(block):
      if sub.get('type')=='table_caption' and key not in self.joins:
       old=self.block;self.block=int(sub.get('index',old));self.para(sub,bold=True);self.block=old
     if decision=='figure':
      if key in self.joins:raise ValueError('Cannot join a source figure as a native table')
      self.picture({'bbox':review.get('bbox') or bodies[0].get('bbox') or block.get('bbox')})
     else:self.table(bodies[0]['content'],join=key in self.joins,cell_images=review.get('cell_images',[]))
     if key in self.joins:
      table_state=(self.last_kind,self.last_table,self.last_cols,self.last_location)
      for sub in children(block):
       if sub.get('type')=='table_caption':
        old=self.block;self.block=int(sub.get('index',old));self.para(sub,bold=True);self.block=old
      self.last_kind,self.last_table,self.last_cols,self.last_location=table_state
     for sub in children(block):
      if sub.get('type')=='table_footnote':self.para(sub)
    elif typ in ('image','chart','figure'):
     self.picture(block)
     for sub in children(block):
      if sub.get('type') in ('image_caption','figure_caption','chart_caption'):self.para(sub)
    elif typ in ('list','ordered_list','unordered_list'):
     for sub in children(block):self.para(sub)
    elif typ=='code':self.code(block)
    elif typ in known_text:
     p=self.para(block,heading=typ=='paragraph_title',title=typ=='doc_title')
    else:
     self.review('unsupported_block_type',type=typ,text=plain(block)[:180])
     if plain(block):self.para(block)
  if toc_page and not inserted:raise ValueError('toc_before_page was excluded or does not exist')
  if self.joins-self.used_joins:raise ValueError(f'Unused confirmed joins: {self.joins-self.used_joins}')
  self.normalize_tables()
  # Source TOCs may use full-width punctuation while body OCR uses ASCII.
  heading_key=lambda text:re.sub(r'\s','',text).replace('；',';').replace('：',':')
  found={heading_key(x['text']) for x in self.report['headings']}
  self.report['source_toc_entries']=len(self.expected)
  self.report['unmatched_source_toc']=[num+' '+title for num,title in self.expected.items() if heading_key(num+title) not in found]
  return self.doc

def main():
 ap=argparse.ArgumentParser()
 for name in ('result','pages','config','reviews','output'):ap.add_argument('--'+name,required=True)
 ap.add_argument('--overwrite',action='store_true');a=ap.parse_args()
 output=Path(a.output).resolve()
 if output.exists() and not a.overwrite:raise SystemExit('Output exists; use a new path or explicit --overwrite')
 data=json.loads(Path(a.result).read_text(encoding='utf-8-sig'));pages=json.loads(Path(a.pages).read_text(encoding='utf-8-sig'));cfg=json.loads(Path(a.config).read_text(encoding='utf-8-sig'))
 cfg['_reviews']=json.loads(Path(a.reviews).read_text(encoding='utf-8-sig'))
 if not cfg.get('title'):raise SystemExit('Read source cover and set title in document.json')
 indexes=[p['page_idx'] for p in data['pages']]
 if indexes!=list(range(len(pages))):raise SystemExit('OCR result page indexes do not match pages.json; do not guess mapping')
 builder=Builder(data,pages,cfg);doc=builder.run();output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
 builder.report['artifact_sha256_at_build']=hashlib.sha256(output.read_bytes()).hexdigest()
 builder.report['inputs_sha256']={key:hashlib.sha256(Path(getattr(a,key)).read_bytes()).hexdigest() for key in ('result','pages','config','reviews')}
 report=output.with_suffix('.build_report.json');report.write_text(json.dumps(builder.report,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps({'output':str(output),'headings':len(builder.report['headings']),'tables':builder.report['tables'],
   'images':builder.report['images'],'review_items':len(builder.report['review_items']),'unmatched_toc':builder.report['unmatched_source_toc']},ensure_ascii=True))

if __name__=='__main__':main()
